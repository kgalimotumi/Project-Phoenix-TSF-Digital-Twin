from __future__ import annotations

import re
from collections import defaultdict
from dataclasses import dataclass, field
from pathlib import Path

JOB_RE = re.compile(r"^(\d{4})")
RN_RE = re.compile(r"\bRN\s*[-_ ]?\s*0*(\d+)\b", re.I)
INVALID = re.compile(r'[<>:"/\\|?*]+')


def normalize_rn(value):
    m = RN_RE.search(str(value or ""))
    return str(int(m.group(1))) if m else ""


def pdf_text(path, pages=5):
    from pypdf import PdfReader
    try:
        reader=PdfReader(str(path)); return "\n".join((p.extract_text() or "") for p in reader.pages[:pages])
    except Exception:
        return ""


def safe_name(text):
    return INVALID.sub("_", str(text)).strip(" ._")[:180] or "Concrete Test Results"


def pdfs(folder, recursive):
    root=Path(folder); return sorted(root.rglob("*.pdf") if recursive else root.glob("*.pdf"),key=lambda p:p.name.lower())


def register_rows(path):
    if not path: return []
    from openpyxl import load_workbook
    wb=load_workbook(path,read_only=True,data_only=True); ws=wb.active
    rows=list(ws.iter_rows(values_only=True));
    if not rows:return []
    header_i=0
    for i,row in enumerate(rows[:25]):
        joined=" | ".join(str(x or "") for x in row).lower()
        if "description" in joined and ("inspection" in joined or "request" in joined or "rn" in joined): header_i=i; break
    headers=[str(x or "").strip().lower() for x in rows[header_i]]
    def col(words):
        return next((i for i,h in enumerate(headers) if any(w in h for w in words)),None)
    ci_rn=col(("request number","inspection number"," rn","rn ","rn")); ci_desc=col(("description","inspection title","title")); ci_job=col(("lab","reference","job"))
    out=[]
    for row in rows[header_i+1:]:
        rn=normalize_rn(row[ci_rn] if ci_rn is not None and ci_rn<len(row) else "")
        desc=str(row[ci_desc] or "").strip() if ci_desc is not None and ci_desc<len(row) else ""
        job=""
        if ci_job is not None and ci_job<len(row):
            m=re.search(r"\b(\d{4})\b",str(row[ci_job] or "")); job=m.group(1) if m else ""
        if rn or desc or job: out.append({"rn":rn,"description":desc,"job":job})
    return out


@dataclass
class Match:
    job: str
    rn: str = ""
    description: str = ""
    cover: list[Path] = field(default_factory=list)
    rn_reports: list[Path] = field(default_factory=list)
    other_lab: list[Path] = field(default_factory=list)
    raw: list[Path] = field(default_factory=list)
    status: str = ""
    output: Path | None = None

    @property
    def lab_order(self): return self.cover+self.rn_reports+self.other_lab
    @property
    def ready(self): return bool(self.rn and self.raw and self.lab_order)


def preview(lab_folder, raw_folder, register_file="", recursive=True):
    labs=pdfs(lab_folder,recursive); raws=pdfs(raw_folder,recursive); groups=defaultdict(list)
    for p in labs:
        m=JOB_RE.match(p.name)
        if m: groups[m.group(1)].append(p)
    register=register_rows(register_file); by_rn={r["rn"]:r for r in register if r["rn"]}; by_job={r["job"]:r for r in register if r["job"]}
    matches=[]
    for job,files in sorted(groups.items()):
        cover=[p for p in files if "cover" in p.name.lower()]; noncover=[p for p in files if p not in cover]
        rn=""
        for p in noncover+cover:
            rn=normalize_rn(p.name)
            if rn: break
        if not rn:
            for p in noncover+cover:
                rn=normalize_rn(pdf_text(p))
                if rn: break
        rn_reports=[p for p in noncover if rn and normalize_rn(p.name)==rn]
        others=[p for p in noncover if p not in rn_reports]
        raw_matches=[]
        for p in raws:
            raw_job=JOB_RE.match(p.name); raw_rn=normalize_rn(p.name)
            if (raw_job and raw_job.group(1)==job) or (rn and raw_rn==rn): raw_matches.append(p)
        reg=by_job.get(job) or by_rn.get(rn,{})
        m=Match(job=job,rn=rn,description=reg.get("description",""),cover=cover,rn_reports=rn_reports,other_lab=others,raw=raw_matches)
        m.status="Ready" if m.ready else ("No RN" if not rn else "Missing raw data")
        matches.append(m)
    known=set(groups); unmatched=[p for p in labs if not JOB_RE.match(p.name)]
    register_missing=[r for r in register if (r["job"] and r["job"] not in known) or (not r["job"] and r["rn"] and not any(m.rn==r["rn"] for m in matches))]
    return matches,unmatched,register_missing


def merge_ready(matches, output_folder, overwrite=False):
    from pypdf import PdfReader, PdfWriter
    out=Path(output_folder); out.mkdir(parents=True,exist_ok=True)
    completed=[]
    for m in matches:
        if not m.ready: continue
        title=m.description or f"RN{m.rn} Concrete Test Results - {m.job}"
        target=out/(safe_name(title)+".pdf"); m.output=target
        if target.exists() and not overwrite:
            m.status="Exists (skipped)"; continue
        writer=PdfWriter()
        try:
            for p in m.lab_order+m.raw:
                reader=PdfReader(str(p))
                for page in reader.pages: writer.add_page(page)
            with open(target,"wb") as f: writer.write(f)
            m.status="Merged"; completed.append(m)
        except Exception as exc: m.status=f"Error: {exc}"
    return completed


def write_report(path, matches, unmatched, register_missing):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb=Workbook(); ws=wb.active; ws.title="Matches"
    headers=["4-digit reference","RN","Status","Inspection description","Coverpage","Lab report(s)","Raw data file(s)","Output filename"]
    ws.append(headers)
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=PatternFill("solid",fgColor="1F4E78")
    fills={"Merged":"C6EFCE","Ready":"C6EFCE","Missing raw data":"FFC7CE","No RN":"FFC7CE","Exists (skipped)":"FFEB9C"}
    for m in matches:
        ws.append([m.job,f"RN{m.rn}" if m.rn else "",m.status,m.description,"; ".join(p.name for p in m.cover),"; ".join(p.name for p in m.rn_reports+m.other_lab),"; ".join(p.name for p in m.raw),m.output.name if m.output else ""])
        colour=fills.get(m.status,"FFC7CE" if m.status.startswith("Error") else "FFFFFF")
        for c in ws[ws.max_row]: c.fill=PatternFill("solid",fgColor=colour)
    miss=wb.create_sheet("Missing inspections"); miss.append(["RN","4-digit reference","Inspection description"])
    for r in register_missing: miss.append([f"RN{r['rn']}" if r['rn'] else "",r["job"],r["description"]])
    un=wb.create_sheet("Unmatched lab files"); un.append(["Filename","Path"])
    for p in unmatched: un.append([p.name,str(p)])
    summary=wb.create_sheet("Summary",0); summary.append(["Metric","Count"]); summary.append(["Lab jobs",len(matches)]); summary.append(["Merged",sum(m.status=="Merged" for m in matches)]); summary.append(["Ready",sum(m.status=="Ready" for m in matches)]); summary.append(["Missing raw data",sum(m.status=="Missing raw data" for m in matches)]); summary.append(["No RN",sum(m.status=="No RN" for m in matches)]); summary.append(["Register inspections missing lab files",len(register_missing)]); summary.append(["Unmatched lab files",len(unmatched)])
    for sheet in wb.worksheets:
        sheet.freeze_panes="A2"; sheet.auto_filter.ref=sheet.dimensions
        for col in sheet.columns:
            letter=col[0].column_letter; sheet.column_dimensions[letter].width=min(60,max(12,max(len(str(c.value or "")) for c in col)+2))
    wb.save(path)

